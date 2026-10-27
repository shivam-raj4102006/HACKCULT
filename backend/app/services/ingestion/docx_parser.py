import os
from typing import List, Dict, Any, Tuple
import docx

def extract_docx_content(file_path: str) -> Tuple[int, List[Dict[str, Any]]]:
    """
    Extracts text, headings, and tables from DOCX documents.
    Returns: (total_sections, list_of_sections)
    """
    if not os.path.exists(file_path):
        return 0, []

    doc = docx.Document(file_path)
    lines = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                lines.append(" | ".join(row_texts))

    full_text = "\n\n".join(lines).strip()
    
    # Chunk long DOCX into logical pages of ~2000 characters if multi-page
    chunk_size = 2000
    if len(full_text) <= chunk_size:
        sections_data = [{
            "page_or_slide_number": 1,
            "text": full_text,
            "character_count": len(full_text),
            "speaker_notes": None
        }]
    else:
        sections_data = []
        chunks = [full_text[i:i+chunk_size] for i in range(0, len(full_text), chunk_size)]
        for idx, chunk in enumerate(chunks):
            sections_data.append({
                "page_or_slide_number": idx + 1,
                "text": chunk.strip(),
                "character_count": len(chunk.strip()),
                "speaker_notes": None
            })

    total_sections = len(sections_data)
    return total_sections, sections_data
