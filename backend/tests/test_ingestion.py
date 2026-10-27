import io
import zipfile
from fastapi.testclient import TestClient
from pptx import Presentation
from docx import Document as DocxDocument
from PIL import Image

def create_sample_pptx_bytes() -> bytes:
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "Rani Ki Vav Archaeological Survey"
    subtitle = slide.placeholders[1]
    subtitle.text = "Commissioned by Queen Udayamati in 1063 CE."
    
    # Add speaker notes
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = "Crucial evidence: The subterranean corridor extends seven levels deep."

    bio = io.BytesIO()
    prs.save(bio)
    return bio.getvalue()

def create_sample_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("Modhera Sun Temple Field Notes", level=1)
    doc.add_paragraph("Constructed in 1026-27 CE during the reign of Bhima I of the Chaulukya dynasty.")
    doc.add_paragraph("The complex consists of Gudhamandapa, Sabhamandapa, and the sacred reservoir Kunda.")
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()

def create_sample_image_bytes() -> bytes:
    img = Image.new("RGB", (100, 100), color=(180, 100, 50))
    bio = io.BytesIO()
    img.save(bio, format="JPEG")
    return bio.getvalue()

def test_file_validation_rejects_unsupported_extensions(client: TestClient):
    fake_exe = io.BytesIO(b"MZ\x90\x00\x03\x00\x00\x00")
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("malicious.exe", fake_exe, "application/x-msdownload")}
    )
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]

def test_pptx_ingestion_and_speaker_notes_extraction(client: TestClient):
    pptx_bytes = create_sample_pptx_bytes()
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("rani_ki_vav.pptx", io.BytesIO(pptx_bytes), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    )
    assert response.status_code == 201
    doc_data = response.json()
    assert doc_data["original_filename"] == "rani_ki_vav.pptx"
    assert doc_data["sha256_hash"] is not None
    doc_id = doc_data["id"]

    # Check detail endpoint
    detail = client.get(f"/api/v1/documents/{doc_id}")
    assert detail.status_code == 200
    detail_data = detail.json()
    assert detail_data["status"] == "processed"
    assert detail_data["total_pages_or_slides"] >= 1
    assert "slide_count" in detail_data["file_metadata"]

def test_docx_ingestion_and_text_extraction(client: TestClient):
    docx_bytes = create_sample_docx_bytes()
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("modhera_notes.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert response.status_code == 201
    doc_id = response.json()["id"]

    detail = client.get(f"/api/v1/documents/{doc_id}")
    assert detail.status_code == 200
    assert detail.json()["status"] == "processed"
    assert detail.json()["total_pages_or_slides"] >= 1

def test_image_ingestion(client: TestClient):
    img_bytes = create_sample_image_bytes()
    response = client.post(
        "/api/v1/documents/upload",
        files={"file": ("pillar_carving.jpg", io.BytesIO(img_bytes), "image/jpeg")}
    )
    assert response.status_code == 201
    doc_id = response.json()["id"]

    detail = client.get(f"/api/v1/documents/{doc_id}")
    assert detail.status_code == 200
    assert detail.json()["status"] == "processed"
    assert detail.json()["file_metadata"]["width"] == 100
    assert detail.json()["file_metadata"]["height"] == 100

def test_document_deduplication_via_sha256(client: TestClient):
    # Unique image to avoid hash collisions with earlier tests
    img = Image.new("RGB", (120, 120), color=(220, 180, 90))
    bio = io.BytesIO()
    img.save(bio, format="JPEG")
    unique_bytes = bio.getvalue()

    res1 = client.post(
        "/api/v1/documents/upload",
        files={"file": ("unique_image_one.jpg", io.BytesIO(unique_bytes), "image/jpeg")}
    )
    res2 = client.post(
        "/api/v1/documents/upload",
        files={"file": ("unique_image_two_identical.jpg", io.BytesIO(unique_bytes), "image/jpeg")}
    )
    assert res1.status_code == 201
    assert res2.status_code == 200  # Returns existing document without creating duplicate
    assert res1.json()["id"] == res2.json()["id"]
    assert res1.json()["sha256_hash"] == res2.json()["sha256_hash"]

def test_zip_archive_batch_extraction(client: TestClient):
    # Create an in-memory ZIP containing a DOCX and a JPEG
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as zf:
        zf.writestr("archive_notes.docx", create_sample_docx_bytes())
        zf.writestr("archive_photo.jpg", create_sample_image_bytes())

    zip_bytes = zip_buffer.getvalue()

    response = client.post(
        "/api/v1/documents/upload-zip",
        files={"file": ("heritage_bundle.zip", io.BytesIO(zip_bytes), "application/zip")}
    )
    assert response.status_code == 201
    batch_data = response.json()
    assert batch_data["total_extracted"] == 2
    assert len(batch_data["documents"]) == 2

def test_zipslip_traversal_attack_blocked(client: TestClient):
    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as zf:
        # Malicious ZipSlip path
        zf.writestr("../../etc/malicious.txt", b"exploit payload")

    response = client.post(
        "/api/v1/documents/upload-zip",
        files={"file": ("exploit.zip", io.BytesIO(zip_buffer.getvalue()), "application/zip")}
    )
    assert response.status_code == 400
    assert "Security violation" in response.json()["detail"]
