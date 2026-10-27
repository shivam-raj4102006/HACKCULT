import io
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session
from pptx import Presentation

from app.models.site import HeritageSite
from app.models.evidence import EvidenceRecord
from app.models.audit import AuditLog
from app.services.seed import seed_gujarat_heritage_data

def test_complete_end_to_end_heritageproof_workflow(client: TestClient, db_session: Session):
    """
    End-to-End Workflow Acceptance Test:
    1. Seed base heritage sites
    2. Register & authenticate human verifier
    3. Ingest multi-format archive (PPTX with speaker notes)
    4. Verify initial unverified status (AI cannot auto-verify)
    5. Human verifier transitions claim to 'verified' with audit trail
    6. Grounded RAG Search returns cited answer referencing the document
    7. Generate & download complete PDF dossier
    """
    # 1. Seed base data
    seed_gujarat_heritage_data(db_session)
    site = db_session.query(HeritageSite).filter(HeritageSite.name == "Rani Ki Vav").first()
    assert site is not None

    # 2. Register & authenticate human verifier
    reg_res = client.post("/api/v1/auth/register", json={
        "username": "dr_patel",
        "email": "dr_patel_archaeology@heritage.org",
        "role": "verifier",
        "password": "SecurePassword123!"
    })
    assert reg_res.status_code == 201

    login_res = client.post("/api/v1/auth/login", json={
        "username": "dr_patel",
        "password": "SecurePassword123!"
    })
    assert login_res.status_code == 200
    token = login_res.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 3. Ingest presentation document
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Rani Ki Vav Archaeological Excavations"
    slide.placeholders[1].text = "Saraswati river silt deposits buried the lower subterranean terraces for centuries."
    notes = slide.notes_slide.notes_text_frame
    notes.text = "ASI desiltation survey conducted between 1986 and 1988 revealed pristine carvings."

    bio = io.BytesIO()
    prs.save(bio)
    pptx_bytes = bio.getvalue()

    upload_res = client.post(
        "/api/v1/documents/upload",
        files={"file": ("excavation_report_1988.pptx", io.BytesIO(pptx_bytes), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    )
    assert upload_res.status_code == 201
    doc_id = upload_res.json()["id"]

    # 4. Check that extracted evidence starts strictly unverified
    ev_records = db_session.query(EvidenceRecord).filter(EvidenceRecord.document_id == doc_id).all()
    assert len(ev_records) > 0
    target_ev = ev_records[0]
    assert target_ev.verification_status == "unverified"
    target_ev.site_id = site.id
    db_session.commit()

    # 5. Human-only verification via authenticated verifier
    review_res = client.post(
        "/api/v1/evidence/review",
        json={
            "evidence_id": target_ev.id,
            "status": "verified",
            "verification_notes": "Corroborated with Archaeological Survey of India 1988 monograph."
        },
        headers=headers
    )
    assert review_res.status_code == 200
    assert review_res.json()["verification_status"] == "verified"

    # Verify audit log was recorded
    audit = db_session.query(AuditLog).filter(AuditLog.entity_id == target_ev.id).first()
    assert audit is not None
    assert audit.action == "CLAIM_VERIFIED"
    assert "Corroborated with Archaeological Survey" in audit.details.get("notes", "")

    # 6. Evidence-Grounded Search on newly verified claim
    qa_res = client.post("/api/v1/search/qa", json={
        "query": "What did the Saraswati river silt deposits bury?",
        "only_verified": True
    })
    assert qa_res.status_code == 200
    qa_data = qa_res.json()
    assert qa_data["status"] == "GROUNDED_ANSWER"
    assert len(qa_data["citations"]) > 0
    assert "excavation_report_1988.pptx" in [c["document_name"] for c in qa_data["citations"]]

    # 7. Generate & download complete PDF dossier
    pdf_res = client.get(f"/api/v1/reports/site/{site.id}/download-direct")
    assert pdf_res.status_code == 200
    assert pdf_res.headers["content-type"] == "application/pdf"
    assert len(pdf_res.content) > 1000
    assert pdf_res.content[:4] == b"%PDF"
